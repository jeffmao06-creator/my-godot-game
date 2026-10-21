#!/usr/bin/env python3
"""Render an editable DOCX and a designed 16:9 PDF from a small slide JSON.

Usage: python3 build_deck.py --content deck_content.json --out-dir . --prefix title
Required: reportlab, python-docx, fonttools. Images are resolved beside the JSON.
No content or generated artwork is altered by this renderer.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import math
import os
from pathlib import Path
from typing import Any

from reportlab.pdfgen import canvas
from reportlab.lib.colors import HexColor, Color
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont as RLTTFont
from reportlab.lib.utils import ImageReader
from fontTools.ttLib import TTFont
from fontTools import subset
from fontTools.fontBuilder import FontBuilder
from fontTools.pens.ttGlyphPen import TTGlyphPen
from fontTools.pens.cu2quPen import Cu2QuPen
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT

W, H = 1280, 720
C = {"paper":"#F4EFE4", "navy":"#17323D", "teal":"#287F7C", "amber":"#D9A648", "red":"#BF5A50", "muted":"#687C80", "line":"#D7DCD4", "white":"#FFFFFF", "lightteal":"#E2ECE7"}
FONT = "DeckTC"
BOLD = "DeckTCBold"
_WARNINGS: list[str] = []


def collect_strings(value: Any) -> str:
    if isinstance(value, str):
        return value
    if isinstance(value, dict):
        return "".join(collect_strings(v) for v in value.values())
    if isinstance(value, list):
        return "".join(collect_strings(v) for v in value)
    return str(value or "")


def register_fonts(content: dict, cache: Path) -> None:
    """Subset CFF Noto TC and convert to glyf, which ReportLab embeds portably."""
    extras = "0123456789ABCDEFGHIJKLMNOPQRSTUVWXYZabcdefghijklmnopqrstuvwxyz ·／｜—–-+×÷%：；，。！？（）「」【】→↑↓<>=/\\.#$&_年月日時星期角色案例金錢疲勞體力能力口才人氣律師受僱合署開業案件員工據點財務排程存讀檔早餐寫書狀查資料跳槽早午後傍晚下班今日收入資金警告示意概念圖介面設計素材製作規格僅供參考原型不代表遊戲畫面NT$臺北本所一二三四五六七八九十生涯第"
    chars = "".join(sorted(set(collect_strings(content)+extras+Path(__file__).read_text(encoding="utf-8"))))
    digest = hashlib.sha256(chars.encode()).hexdigest()[:16]
    cache.mkdir(parents=True, exist_ok=True)
    for name, weight in [(FONT,"Regular"),(BOLD,"Bold")]:
        target = cache / f"noto-tc-{weight.lower()}-{digest}.ttf"
        if not target.exists():
            source = Path(f"/usr/share/fonts/opentype/noto/NotoSansCJK-{weight}.ttc")
            if not source.exists():
                raise FileNotFoundError("Traditional Chinese Noto CJK fonts are required")
            old = TTFont(source, fontNumber=3)
            options = subset.Options()
            options.layout_features = []
            options.hinting = False
            options.desubroutinize = True
            sub = subset.Subsetter(options=options)
            sub.populate(text=chars)
            sub.subset(old)
            glyphset = old.getGlyphSet()
            glyphs = {}
            order = old.getGlyphOrder()
            for glyph_name in order:
                pen = TTGlyphPen(glyphset)
                quad = Cu2QuPen(pen, max_err=1.0, reverse_direction=True)
                glyphset[glyph_name].draw(quad)
                glyphs[glyph_name] = pen.glyph()
            fb = FontBuilder(old["head"].unitsPerEm, isTTF=True)
            fb.setupGlyphOrder(order)
            fb.setupCharacterMap(old.getBestCmap())
            fb.setupGlyf(glyphs)
            fb.setupHorizontalMetrics({g:old["hmtx"].metrics[g] for g in order})
            fb.setupHorizontalHeader(ascent=old["hhea"].ascent, descent=old["hhea"].descent, lineGap=0)
            family = f"LawFirmDeckTC{weight}"
            fb.setupNameTable({"familyName":family,"styleName":weight,"uniqueFontIdentifier":family+digest,"fullName":family,"psName":family,"version":"Version 1.0"})
            os2 = old["OS/2"]
            fb.setupOS2(sTypoAscender=os2.sTypoAscender,sTypoDescender=os2.sTypoDescender,sTypoLineGap=0,usWinAscent=os2.usWinAscent,usWinDescent=os2.usWinDescent)
            fb.setupPost()
            fb.save(target)
            old.close()
        pdfmetrics.registerFont(RLTTFont(name,str(target)))


def color(name: str):
    return HexColor(C.get(name,name))


def box(c,x,y,w,h,fill="white",stroke=None,r=0):
    c.setFillColor(color(fill))
    c.setStrokeColor(color(stroke or fill))
    if r:
        c.roundRect(x,y,w,h,r,fill=1,stroke=bool(stroke))
    else:
        c.rect(x,y,w,h,fill=1,stroke=bool(stroke))


def lines_for(text: Any, width: float, size: float, font: str=FONT):
    text = str(text or "")
    out=[]
    for para in text.split("\n"):
        if not para:
            out.append("")
            continue
        line=""
        for char in para:
            if line and pdfmetrics.stringWidth(line+char,font,size)>width:
                out.append(line.rstrip())
                line=char.lstrip() if char.isspace() else char
            else:
                line+=char
        out.append(line.rstrip())
    return out


def textbox(c,text,x,top,w,h,size=22,fill="navy",font=FONT,leading=None,min_size=16,align="left"):
    if text is None:
        return 0
    original=size
    while True:
        ld=leading or size*1.48
        rows=lines_for(text,w,size,font)
        need=len(rows)*ld
        if need<=h+0.01 or size<=min_size:
            break
        size-=0.5
    if need>h+0.01:
        raise ValueError(f"Text overflow {need:.1f}>{h:.1f}: {str(text)[:90]}")
    if size<original-3:
        _WARNINGS.append(f"Text fitted {original}→{size}: {str(text)[:50]}")
    c.setFont(font,size)
    c.setFillColor(color(fill))
    y=top-size
    for row in rows:
        if align=="center":
            c.drawCentredString(x+w/2,y,row)
        elif align=="right":
            c.drawRightString(x+w,y,row)
        else:
            c.drawString(x,y,row)
        y-=ld
    return need


def bullet_list(c,items,x,top,w,h,size=22,fill="navy"):
    if not items:
        return
    # Reserve proportionally by actual wrapped line count, then fit once globally.
    size0=size
    while size>=16:
        total=sum(len(lines_for(str(item),w-28,size))*size*1.48+18 for item in items)
        if total<=h:
            break
        size-=.5
    if size<16:
        raise ValueError("Bullet list too long: "+str(items)[:100])
    y=top
    for item in items:
        box(c,x,y-size+2,8,8,"amber")
        used=textbox(c,str(item),x+28,y,w-28,h,size=size,fill=fill,min_size=size)
        y-=used+18


def resolve_image(slide: dict, base: Path):
    image=slide.get("image")
    if not image:
        return None
    path=Path(image)
    if not path.is_absolute():
        path=base/path
    if not path.exists():
        raise FileNotFoundError(path)
    return path


def draw_image(c,path,x,y,w,h,background="white",border=False):
    if background:
        box(c,x,y,w,h,background,stroke="line" if border else None,r=0)
    if not path:
        textbox(c,"概念圖",x,y+h/2+20,w,60,size=28,fill="muted",align="center")
        return
    reader=ImageReader(str(path))
    iw,ih=reader.getSize()
    scale=min(w/iw,h/ih)
    dw,dh=iw*scale,ih*scale
    c.drawImage(reader,x+(w-dw)/2,y+(h-dh)/2,width=dw,height=dh,mask="auto")


def footer(c,index,total,slide,dark=False):
    col="paper" if dark else "muted"
    c.setStrokeColor(color("teal" if dark else "line"))
    c.setLineWidth(.8)
    c.line(56,51,W-56,51)
    label="LAW FIRM SIMULATOR  /  ART & GAME DESIGN"
    textbox(c,label,56,37,900,20,size=11,fill=col,min_size=11)
    textbox(c,f"{index:02d} / {total:02d}",W-140,38,84,22,size=13,fill=col,min_size=13,align="right")
    note=slide.get("footnote")
    if note:
        textbox(c,note,56,79,1168,23,size=12,fill=col,min_size=10)


def header(c,slide,index,total):
    box(c,0,0,W,H,"paper")
    for dx,dy,s,co in [(0,0,14,"amber"),(19,0,14,"teal"),(38,0,14,"navy"),(38,19,14,"teal")]:
        box(c,W-109+dx,H-67+dy,s,s,co)
    kicker=slide.get("kicker", "GAME DESIGN / ART DIRECTION")
    textbox(c,kicker,56,H-40,1080,24,size=13,fill="teal",font=BOLD,min_size=13)
    textbox(c,slide.get("title",""),56,H-84,1168,66,size=38,font=BOLD,min_size=30)
    if slide.get("subtitle"):
        textbox(c,slide["subtitle"],58,H-158,1150,50,size=19,fill="muted",min_size=17)
    footer(c,index,total,slide)


def draw_cards(c,cards,x=56,y=105,w=1168,h=405,columns=None):
    if not cards:
        return
    n=len(cards)
    cols=columns or (2 if n in (2,4) else 3)
    rows=math.ceil(n/cols)
    gap=20
    cw=(w-gap*(cols-1))/cols
    ch=(h-gap*(rows-1))/rows
    accents=["teal","amber","navy","red"]
    for i,item in enumerate(cards):
        col,row=i%cols,i//cols
        xx=x+col*(cw+gap); yy=y+h-(row+1)*ch-row*gap
        box(c,xx,yy,cw,ch,"white",r=6)
        box(c,xx,yy+ch-6,cw,6,accents[i%len(accents)])
        tag=item.get("tag")
        title_top=yy+ch-24
        if tag:
            textbox(c,tag,xx+22,title_top,cw-44,25,size=13,fill=accents[i%len(accents)],font=BOLD,min_size=11)
            title_top-=30
        title=str(item.get("title", ""))
        ts=26 if cw>450 else 25
        title_h=70 if len(title)>10 else 45
        used=textbox(c,title,xx+22,title_top,cw-44,title_h,size=ts,font=BOLD,min_size=20)
        body=item.get("body","")
        if isinstance(body,list): body="\n".join(body)
        remaining=title_top-used-16-(yy+20)
        textbox(c,body,xx+22,title_top-used-15,cw-44,remaining,size=21 if rows==1 else 19,min_size=16)


def draw_table(c,table,x=56,top=510,w=1168,h=405):
    headers=table.get("headers",[]); rows=table.get("rows",[])
    if not headers:
        return
    n=len(headers)
    raw=table.get("widths") or [1]*n
    widths=[w*a/sum(raw) for a in raw]
    fs=20
    while fs>=15:
        hh=max(len(lines_for(t,widths[i]-28,fs,BOLD)) for i,t in enumerate(headers))*fs*1.35+24
        rhs=[max(len(lines_for(str(cell),widths[i]-28,fs)) for i,cell in enumerate(row))*fs*1.4+24 for row in rows]
        if hh+sum(rhs)<=h:
            break
        fs-=.5
    if fs<15:
        raise ValueError("Table too tall; reduce rows/content")
    y=top
    xx=x
    for i,text in enumerate(headers):
        box(c,xx,y-hh,widths[i],hh,"navy")
        textbox(c,text,xx+14,y-12,widths[i]-28,hh-18,size=fs,font=BOLD,fill="white",min_size=fs)
        xx+=widths[i]
    y-=hh
    for ri,row in enumerate(rows):
        rh=rhs[ri]; xx=x
        for i,cell in enumerate(row):
            box(c,xx,y-rh,widths[i],rh,"white" if ri%2==0 else "lightteal")
            textbox(c,str(cell),xx+14,y-12,widths[i]-28,rh-18,size=fs,min_size=fs)
            xx+=widths[i]
        y-=rh


def draw_flow(c,cards):
    n=len(cards)
    if not n: return
    gap=32
    cw=(1168-gap*(n-1))/n
    for i,item in enumerate(cards):
        x=56+i*(cw+gap)
        box(c,x,175,cw,315,"white",r=6)
        box(c,x,430,cw,60,"teal" if i%2==0 else "navy")
        textbox(c,str(item.get("tag") or f"{i+1:02d}"),x+20,475,cw-40,40,size=22,fill="white",font=BOLD,min_size=18)
        textbox(c,item.get("title",""),x+20,405,cw-40,75,size=27,font=BOLD,min_size=20)
        textbox(c,item.get("body",""),x+20,320,cw-40,120,size=20,min_size=16)
        if i<n-1:
            mid=x+cw+gap/2
            c.setFillColor(color("amber"))
            p=c.beginPath();p.moveTo(mid-8,338);p.lineTo(mid+8,329);p.lineTo(mid-8,320);p.close()
            c.drawPath(p,fill=1,stroke=0)


def draw_ui(c,slide,base):
    x,y,w,h=56,105,900,405
    box(c,x,y,w,h,"navy",r=5)
    box(c,x+8,y+h-82,w-16,74,"paper")
    textbox(c,"生涯第1年 02月03日｜09:00",x+22,y+h-22,510,28,size=17,font=BOLD,min_size=17)
    textbox(c,"資金 NT$50,000",x+560,y+h-22,300,28,size=17,font=BOLD,min_size=17)
    textbox(c,"疲勞 85.6/100",x+22,y+h-53,250,25,size=16,font=BOLD,min_size=16)
    textbox(c,"在手 16 案",x+300,y+h-53,200,25,size=16,font=BOLD,min_size=16)
    textbox(c,"明早負荷 14.4",x+560,y+h-53,300,25,size=16,font=BOLD,min_size=16)
    box(c,x+8,y+45,665,h-132,"lightteal")
    draw_image(c,resolve_image(slide,base),x+8,y+45,665,h-132,background="lightteal")
    box(c,x+681,y+45,211,h-132,"white")
    textbox(c,"早上行動",x+697,y+h-101,180,35,size=21,font=BOLD,min_size=20)
    for i,t in enumerate(["吃早餐  NT$100","寫書狀","摸魚","跳槽"]):
        yy=y+h-166-i*47
        box(c,x+696,yy,181,38,"teal" if i==1 else "paper",r=3)
        textbox(c,t,x+708,yy+30,157,30,size=16,fill="white" if i==1 else "navy",min_size=16)
    textbox(c,"案件｜角色｜職涯｜員工｜據點｜財務｜排程｜存讀檔",x+20,y+30,w-40,26,size=17,fill="paper",min_size=15)
    box(c,980,105,244,405,"white",r=6)
    box(c,980,504,244,6,"teal")
    textbox(c,"介面需求",1000,481,204,40,size=25,font=BOLD,min_size=22)
    notes=slide.get("bullets") or ["行動前預覽費用與消耗。","案件、員工、財務快速切換。","二十個存檔槽位共用元件。"]
    bullet_list(c,notes,1000,421,204,285,size=20)


def draw_palette(c,slide,base):
    palette=slide.get("palette") or [{"name":"紙張米白","hex":C["paper"],"use":"底色"},{"name":"墨藍","hex":C["navy"],"use":"文字／制服"},{"name":"玉綠","hex":C["teal"],"use":"互動／櫃體"},{"name":"暖金","hex":C["amber"],"use":"燈光／里程碑"},{"name":"危機紅","hex":C["red"],"use":"警告"}]
    leftw=430
    step=min(76,390/max(1,len(palette)))
    for i,item in enumerate(palette):
        yy=510-(i+1)*step
        box(c,56,yy,70,step-12,item.get("hex",C["teal"]),stroke="line",r=3)
        textbox(c,item.get("name",""),146,yy+step-12,330,27,size=21,font=BOLD,min_size=18)
        textbox(c,f'{item.get("hex","")}  {item.get("use","")}',146,yy+23,330,25,size=14.5,fill="muted",min_size=13)
    image=resolve_image(slide,base)
    if image:
        draw_image(c,image,520,120,704,396,background="white",border=True)
    else:
        box(c,520,120,704,390,"white",r=6)
        bullet_list(c,slide.get("bullets",[]),550,475,644,245,size=25)
        # Material swatches are diagrammatic blocks, not modified concept artwork.
        for row in range(3):
            for col in range(12):
                co=["navy","teal","amber","paper"][min(3,(col//3+row)%4)]
                box(c,550+col*52,146+row*24,48,20,co)


def render_pdf(content,source,out):
    c=canvas.Canvas(str(out),pagesize=(W,H),pageCompression=1)
    c.setTitle(content.get("title","律師事務所模擬器｜概念與美術簡報"))
    c.setAuthor("Law Firm Simulator / Design")
    slides=content["slides"]
    for idx,slide in enumerate(slides,1):
        layout=slide.get("layout","cards")
        if layout=="cover":
            box(c,0,0,W,H,"navy")
            image=resolve_image(slide,source.parent)
            draw_image(c,image,475,140,750,422,background="navy")
            textbox(c,slide.get("kicker","LAW FIRM SIMULATOR"),56,656,1100,30,size=15,fill="amber",font=BOLD,min_size=14)
            textbox(c,slide.get("title",""),56,561,386,226,size=57,fill="paper",font=BOLD,min_size=38,leading=78)
            textbox(c,slide.get("subtitle",""),58,310,375,126,size=22,fill="paper",min_size=18)
            for i,b in enumerate(slide.get("bullets",[])[:3]):
                textbox(c,b,58,158-i*29,370,28,size=15,fill="paper",min_size=13)
            footer(c,idx,len(slides),slide,dark=True)
        else:
            header(c,slide,idx,len(slides))
            if layout=="image_right":
                bullet_list(c,slide.get("bullets",[]),58,505,405,390,size=24)
                draw_image(c,resolve_image(slide,source.parent),500,120,724,407,background="white",border=True)
            elif layout=="image_full":
                draw_image(c,resolve_image(slide,source.parent),56,105,1168,410,background="white",border=True)
                if slide.get("image_caption"):
                    textbox(c,slide["image_caption"],70,104,1140,24,size=13,fill="muted",min_size=12,align="center")
            elif layout=="table":
                if slide.get("bullets"):
                    draw_table(c,slide.get("table",{}),h=300)
                    notes=slide["bullets"]
                    for ni,note in enumerate(notes):
                        bullet_list(c,[note],60+ni*(1160/len(notes)),185,1140/len(notes),70,size=20)
                else:
                    draw_table(c,slide.get("table",{}))
            elif layout=="flow":
                draw_flow(c,slide.get("cards",[]))
            elif layout=="palette":
                draw_palette(c,slide,source.parent)
            elif layout=="ui_mock":
                draw_ui(c,slide,source.parent)
            elif layout=="split":
                draw_cards(c,slide.get("cards",[]),columns=2)
            else:
                cards=slide.get("cards",[])
                if cards:
                    draw_cards(c,cards)
                else:
                    bullet_list(c,slide.get("bullets",[]),60,510,1140,400,size=28)
        c.showPage()
    c.save()


def set_run_font(run,size=None,bold=None,color_hex=None):
    run.font.name="Noto Sans CJK TC"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"),"Noto Sans CJK TC")
    if size is not None: run.font.size=Pt(size)
    if bold is not None: run.bold=bold
    if color_hex: run.font.color.rgb=RGBColor.from_string(color_hex.lstrip("#"))


def shade_cell(cell,fill):
    tcPr=cell._tc.get_or_add_tcPr()
    shd=OxmlElement("w:shd");shd.set(qn("w:fill"),fill.lstrip("#"));tcPr.append(shd)


def para(doc,text,size=13,bold=False,color_hex=C["navy"],before=0,after=6):
    p=doc.add_paragraph()
    p.paragraph_format.space_before=Pt(before);p.paragraph_format.space_after=Pt(after)
    p.paragraph_format.line_spacing=1.12
    run=p.add_run(str(text or ""));set_run_font(run,size,bold,color_hex)
    return p


def docx_table(doc,headers,rows,widths=None,fontsize=11.5):
    table=doc.add_table(rows=1,cols=len(headers))
    table.alignment=WD_TABLE_ALIGNMENT.CENTER
    table.autofit=False
    for i,title in enumerate(headers):
        cell=table.rows[0].cells[i];shade_cell(cell,C["navy"])
        cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(6);p.paragraph_format.space_before=Pt(6)
        set_run_font(p.add_run(title),fontsize,True,C["white"])
    for ri,row in enumerate(rows):
        cells=table.add_row().cells
        for i,text in enumerate(row):
            shade_cell(cells[i],C["white"] if ri%2==0 else C["lightteal"])
            p=cells[i].paragraphs[0];p.paragraph_format.space_after=Pt(6);p.paragraph_format.space_before=Pt(6)
            p.paragraph_format.line_spacing=1.08
            set_run_font(p.add_run(str(text)),fontsize,False,C["navy"])
    if widths:
        total=sum(widths)
        for row in table.rows:
            for i,cell in enumerate(row.cells):cell.width=Inches(12.13*widths[i]/total)
    return table


def docx_cards(doc,cards,layout):
    n=len(cards)
    cols=4 if layout=="flow" and n==4 else (3 if n==3 else 2)
    table=doc.add_table(rows=math.ceil(n/cols),cols=cols)
    table.alignment=WD_TABLE_ALIGNMENT.CENTER
    table.autofit=False
    for i,cell in enumerate([cc for row in table.rows for cc in row.cells]):
        cell.width=Inches(12.1/cols)
        if i>=n:continue
        card=cards[i]
        shade_cell(cell,C['white'])
        cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
        p=cell.paragraphs[0]
        p.paragraph_format.space_before=Pt(9);p.paragraph_format.space_after=Pt(7)
        if card.get('tag'):
            set_run_font(p.add_run(card['tag']+"\n"),9,True,C['teal'])
        set_run_font(p.add_run(card.get('title','')),15,True,C['navy'])
        p=cell.add_paragraph()
        p.paragraph_format.space_after=Pt(10)
        p.paragraph_format.line_spacing=Pt(18)
        body=card.get('body','')
        if isinstance(body,list):body='\n'.join(body)
        set_run_font(p.add_run(body),12,False,C['navy'])
    return table


def render_docx(content,source,out):
    doc=Document()
    section=doc.sections[0]
    section.page_width=Inches(13.333);section.page_height=Inches(7.5)
    section.top_margin=Inches(.42);section.bottom_margin=Inches(.43)
    section.left_margin=Inches(.58);section.right_margin=Inches(.58)
    section.header_distance=Inches(.13);section.footer_distance=Inches(.13)
    for grid in list(section._sectPr.findall(qn('w:docGrid'))):section._sectPr.remove(grid)
    normal=doc.styles['Normal']
    normal.font.name="Noto Sans CJK TC";normal.font.size=Pt(12)
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"),"Noto Sans CJK TC")
    normal.paragraph_format.space_after=Pt(5)
    # Set page background as a light paper tone in capable Word processors.
    bg=OxmlElement('w:background');bg.set(qn('w:color'),C["paper"].lstrip('#'));doc._element.insert(0,bg)
    footer_p=section.footer.paragraphs[0]
    footer_p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(footer_p.add_run("LAW FIRM SIMULATOR  /  "),8,color_hex=C["muted"])
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer_p._p.append(fld)
    slides=content['slides']
    for idx,slide in enumerate(slides):
        kicker_para=para(doc,slide.get('kicker',"GAME DESIGN / ART DIRECTION"),10,True,C['teal'],after=3)
        if idx:kicker_para.paragraph_format.page_break_before=True
        para(doc,slide.get('title',''),27,True,after=6)
        if slide.get('subtitle'):para(doc,slide['subtitle'],12.5,color_hex=C['muted'],after=9)
        layout=slide.get('layout','cards')
        image=resolve_image(slide,source.parent)
        handled_bullets=False
        if layout in ('cover','image_right') and image:
            # Side-by-side content preserves a full concept image without squeezing
            # bullets beneath it or creating an accidental spill/blank page.
            tbl=doc.add_table(rows=1,cols=2)
            tbl.alignment=WD_TABLE_ALIGNMENT.CENTER
            tbl.autofit=False
            left,right=tbl.rows[0].cells
            left.width=Inches(4.65);right.width=Inches(7.45)
            for cell in (left,right):cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for bi,text in enumerate(slide.get('bullets',[])):
                pp=left.paragraphs[0] if bi==0 else left.add_paragraph()
                pp.paragraph_format.space_after=Pt(13)
                pp.paragraph_format.line_spacing=Pt(20)
                set_run_font(pp.add_run('▪ '+text),13,False,C['navy'])
            pp=right.paragraphs[0]
            pp.alignment=WD_ALIGN_PARAGRAPH.CENTER
            pp.paragraph_format.space_after=Pt(0)
            from PIL import Image
            with Image.open(image) as im:side_ratio=im.width/im.height
            pp.add_run().add_picture(str(image),width=Inches(min(7.25,3.8*side_ratio)))
            handled_bullets=True
        elif layout in ('image_full','palette','ui_mock') and image:
            from PIL import Image
            with Image.open(image) as im:ratio=im.width/im.height
            max_h=4.0 if layout=='image_full' else 2.9
            max_w=11.9
            ww=min(max_w,max_h*ratio)
            p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after=Pt(6)
            p.add_run().add_picture(str(image),width=Inches(ww))
            if slide.get('image_caption'):para(doc,slide['image_caption'],9.5,color_hex=C['muted'],after=6)
        if slide.get('bullets') and not handled_bullets:
            for text in slide['bullets']:
                para(doc,'▪ '+str(text),12 if image else 15,after=6)
        if slide.get('cards'):
            docx_cards(doc,slide['cards'],layout)
        if slide.get('table'):
            t=slide['table'];docx_table(doc,t.get('headers',[]),t.get('rows',[]),widths=t.get('widths'),fontsize=11)
        if slide.get('palette'):
            docx_table(doc,["色彩","色碼","用途"],[[i.get('name',''),i.get('hex',''),i.get('use','')] for i in slide['palette']],fontsize=10.5)
        if slide.get('footnote'):para(doc,slide['footnote'],9,color_hex=C['muted'],before=5,after=0)
    props=doc.core_properties
    props.title=content.get('title','律師事務所模擬器｜概念與美術簡報')
    props.subject='2D pixel art game design and art production presentation'
    props.author='Law Firm Simulator / Design'
    doc.save(out)


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--content',required=True,type=Path)
    parser.add_argument('--out-dir',type=Path,default=Path('.'))
    parser.add_argument('--prefix',default='Law_Firm_Simulator_Art_Pitch.zh-TW')
    args=parser.parse_args()
    content=json.loads(args.content.read_text(encoding='utf-8'))
    if not isinstance(content.get('slides'),list) or not content['slides']:
        raise ValueError('content.slides must be a nonempty list')
    args.out_dir.mkdir(parents=True,exist_ok=True)
    cache=Path(os.environ.get('TMPDIR','/tmp'))/'lawfirm-deck-fonts'
    register_fonts(content,cache)
    pdf=args.out_dir/(args.prefix+'.pdf');docx=args.out_dir/(args.prefix+'.docx')
    render_pdf(content,args.content,pdf)
    render_docx(content,args.content,docx)
    print(json.dumps({'slides':len(content['slides']),'pdf':str(pdf),'docx':str(docx),'warnings':_WARNINGS},ensure_ascii=False))


if __name__=='__main__':
    main()
