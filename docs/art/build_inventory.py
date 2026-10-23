#!/usr/bin/env python3
"""Build editable inventory Word/PDF from the tracked Markdown.

Usage: python docs/art/build_inventory.py [--docx-only] [--output-dir DIR]
Dependencies: python-docx, LibreOffice (soffice), Noto Sans CJK TC font.
The source Markdown and any other project files are never modified.
"""
from pathlib import Path
import re,subprocess,argparse,tempfile,hashlib,shutil
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT

HERE=Path(__file__).resolve().parent
parser=argparse.ArgumentParser(description='從最終Markdown重建美術製作清單Word與PDF。需要python-docx；PDF另需LibreOffice與Noto Sans CJK TC。')
parser.add_argument('--source',type=Path,default=HERE/'ART_ASSET_LIST.zh-TW.md',help='輸入Markdown；預設為同目錄清單')
parser.add_argument('--output-dir',type=Path,default=HERE,help='Word/PDF輸出目錄')
parser.add_argument('--docx-only',action='store_true',help='只建立Word，不呼叫LibreOffice')
args=parser.parse_args()
source=args.source.resolve()
ART=args.output_dir.resolve();ART.mkdir(parents=True,exist_ok=True)
DOCX=ART/'Art_Production_Inventory.zh-TW.docx'
source_bytes=source.read_bytes()
source_sha=hashlib.sha256(source_bytes).hexdigest()
FONT='Noto Sans CJK TC'
doc=Document()
sec=doc.sections[0]
sec.page_width=Cm(29.7);sec.page_height=Cm(21)
sec.left_margin=Cm(1.35);sec.right_margin=Cm(1.35)
sec.top_margin=Cm(1.25);sec.bottom_margin=Cm(1.25)
sec.header_distance=Cm(.5);sec.footer_distance=Cm(.5)

def fontstyle(style,size,color=None,bold=False):
    f=style.font;f.name=FONT;f.size=Pt(size);f.bold=bold
    if color:f.color.rgb=RGBColor.from_string(color)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),FONT)
fontstyle(doc.styles['Normal'],9.5,'17323D')
doc.styles['Normal'].paragraph_format.space_after=Pt(5)
doc.styles['Normal'].paragraph_format.line_spacing=1.12
for name,size in [('Title',26),('Heading 1',16),('Heading 2',12),('Heading 3',10.5)]:
    fontstyle(doc.styles[name],size,'17323D',True)
    doc.styles[name].paragraph_format.keep_with_next=True
    doc.styles[name].paragraph_format.space_before=Pt(10)
    doc.styles[name].paragraph_format.space_after=Pt(5)

hp=sec.header.paragraphs[0];hp.text='訴訟之神之路  /  美術設定與製作資源清單 v0.2'
for r in hp.runs:r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string('287F7C')
fp=sec.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
fp.add_run('概念／製作規格・正式遊戲素材仍待製作  |  ')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');fp._p.append(fld)
for r in fp.runs:r.font.size=Pt(7.5)

def inline(paragraph,text):
    for part in re.split(r'(\*\*.*?\*\*|`[^`]+`)',text):
        if not part:continue
        if part.startswith('**'):
            run=paragraph.add_run(part[2:-2]);run.bold=True
        elif part.startswith('`'):
            run=paragraph.add_run(part[1:-1]);run.font.color.rgb=RGBColor.from_string('287F7C')
        else:paragraph.add_run(part)

def table(block):
    data=[[v.strip() for v in line.strip().strip('|').split('|')] for line in block]
    data=[row for row in data if not all(re.fullmatch(r':?-+:?',v) for v in row)]
    columns=len(data[0]);t=doc.add_table(rows=1,cols=columns)
    t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    widths=[2.0,5.4,4.1,7,6.5,1.3] if columns==6 else ([3.2,4.3,18.8] if columns==3 else [26.3/columns]*columns)
    for i,w in enumerate(widths):t.columns[i].width=Cm(w)
    for index,values in enumerate(data):
        cells=t.rows[0].cells if index==0 else t.add_row().cells
        for col,(cell,value) in enumerate(zip(cells,values)):
            cell.width=Cm(widths[col]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
            p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3);p.paragraph_format.line_spacing=1.04
            inline(p,value)
            for r in p.runs:r.font.size=Pt(8.1 if columns==6 else 9)
            tcPr=cell._tc.get_or_add_tcPr()
            shading=OxmlElement('w:shd');shading.set(qn('w:fill'),'17323D' if index==0 else ('F4EFE4' if index%2 else 'FFFFFF'));tcPr.append(shading)
            if index==0:
                for r in p.runs:r.bold=True;r.font.color.rgb=RGBColor(255,255,255)
        trPr=t.rows[index]._tr.get_or_add_trPr()
        noSplit=OxmlElement('w:cantSplit');trPr.append(noSplit)
        if index==0:
            repeat=OxmlElement('w:tblHeader');trPr.append(repeat)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

lines=source_bytes.decode('utf-8').splitlines()
i=0
while i<len(lines):
    text=lines[i].strip()
    if not text:i+=1;continue
    if text.startswith('|'):
        block=[]
        while i<len(lines) and lines[i].strip().startswith('|'):block.append(lines[i]);i+=1
        table(block);continue
    if text.startswith('# '):
        doc.add_heading(text[2:].replace('｜',' / '),0)
    elif text.startswith('## '):
        if text.startswith('## 4.'):doc.add_page_break()
        doc.add_heading(text[3:],1)
    elif text.startswith('### '):doc.add_heading(text[4:],2)
    elif text.startswith('- '):inline(doc.add_paragraph(style='List Bullet'),text[2:])
    elif re.match(r'^\d\. ',text):inline(doc.add_paragraph(style='List Number'),text[3:])
    else:inline(doc.add_paragraph(),text)
    i+=1

doc.core_properties.title='訴訟之神之路｜美術設定與製作資源清單 v0.2'
doc.core_properties.subject='分層人物、正交場景與全資源派工清單'
doc.core_properties.author='訴訟之神之路'
doc.core_properties.comments='Source '+source.name+' SHA256 '+source_sha
doc.save(DOCX)
if not args.docx_only:
    soffice=shutil.which('soffice')
    if not soffice:raise SystemExit('Word已產生；找不到soffice。請安裝LibreOffice或使用--docx-only。')
    with tempfile.TemporaryDirectory(prefix='law-art-inventory-') as profile:
        subprocess.run([soffice,'-env:UserInstallation='+Path(profile).as_uri(),'--headless','--convert-to','pdf','--outdir',str(ART),str(DOCX)],check=True)
print(DOCX)
print('Source SHA256 '+source_sha)
